#!/usr/bin/env python3
"""
AI 互联接口综合研究报告 — 合并版 PPTX 生成脚本

融合以下四个 LaTeX 源文件的内容:
  1. llm_interconnect_product_template.tex  — LLM 互联接口产品线调研
  2. ai_interface_research.tex              — AI 领域的接口互联与研究方向
  3. cxl_presentation.tex                   — CXL 调研报告
  4. xconn_presentation.tex                 — XConn Technologies 产品介绍

用法:
    python pptx_merged.py                          → 输出 interconnect_merged.pptx
    python pptx_merged.py -o final_report.pptx

依赖:
    pip install python-pptx
"""

import argparse
import sys
import datetime
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.shapes import MSO_SHAPE
except ImportError:
    print("错误: 需要 python-pptx 库，请运行: pip install python-pptx")
    sys.exit(1)

# ═══════════════════════ 配色方案 ═══════════════════════
COLOR_IE_BLUE   = RGBColor(0, 83, 159)
COLOR_IE_RED    = RGBColor(198, 38, 46)
COLOR_IE_GREEN  = RGBColor(0, 114, 41)
COLOR_IE_ORANGE = RGBColor(235, 110, 31)
COLOR_DARK_GRAY = RGBColor(66, 66, 66)
COLOR_WHITE     = RGBColor(255, 255, 255)
COLOR_BLUE_LIGHT_BG  = RGBColor(220, 235, 250)
COLOR_GREEN_LIGHT_BG = RGBColor(220, 245, 225)
COLOR_RED_LIGHT_BG   = RGBColor(250, 220, 220)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

# ═══════════════════════ 工具函数 ═══════════════════════

def set_slide_bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color

def add_shape(slide, left, top, width, height, fill_color=None, line_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.line.fill.background()
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(1)
    return shape

def add_textbox(slide, left, top, width, height, text, font_size=14,
                font_color=COLOR_DARK_GRAY, bold=False, alignment=PP_ALIGN.LEFT,
                font_name="Microsoft YaHei"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = alignment
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.color.rgb = font_color
    run.font.bold = bold
    run.font.name = font_name
    return txBox

def add_paragraph(tf, text, font_size=13, font_color=COLOR_DARK_GRAY,
                  bold=False, alignment=PP_ALIGN.LEFT, space_before=Pt(4)):
    p = tf.add_paragraph()
    p.alignment = alignment
    p.space_before = space_before
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.color.rgb = font_color
    run.font.bold = bold
    run.font.name = "Microsoft YaHei"

def add_bullet(tf, text, level=0, font_size=12, font_color=COLOR_DARK_GRAY,
               bold=False, bullet_char="•"):
    p = tf.add_paragraph()
    p.level = level
    p.space_before = Pt(2)
    p.space_after = Pt(1)
    run = p.add_run()
    prefix = f"{bullet_char} " if level == 0 else f"  {bullet_char} "
    run.text = f"{prefix}{text}"
    run.font.size = Pt(font_size)
    run.font.color.rgb = font_color
    run.font.bold = bold
    run.font.name = "Microsoft YaHei"

def add_footer_bar(slide):
    add_shape(slide, Inches(0), Inches(7.2), SLIDE_W, Inches(0.3),
              fill_color=COLOR_IE_BLUE)

def add_top_bar(slide):
    add_shape(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.06),
              fill_color=COLOR_IE_BLUE)

# ═══════════════════════ 幻灯片构建 ═══════════════════════

def make_title_slide(prs, title, subtitle="", author=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, COLOR_WHITE)
    add_shape(slide, Inches(0), Inches(0), SLIDE_W, Inches(3.2), fill_color=COLOR_IE_BLUE)
    add_shape(slide, Inches(0), Inches(3.2), SLIDE_W, Inches(0.08), fill_color=COLOR_IE_ORANGE)
    add_textbox(slide, Inches(0.8), Inches(0.6), Inches(11.5), Inches(1.5),
                title, font_size=32, font_color=COLOR_WHITE, bold=True)
    if subtitle:
        add_textbox(slide, Inches(0.8), Inches(2.0), Inches(11.5), Inches(0.8),
                    subtitle, font_size=18, font_color=RGBColor(200, 220, 240))
    today = datetime.date.today().strftime("%Y年%m月%d日")
    add_textbox(slide, Inches(0.8), Inches(4.0), Inches(5), Inches(0.5),
                f"汇报人：{author}" if author else "", font_size=14, font_color=COLOR_DARK_GRAY)
    add_textbox(slide, Inches(0.8), Inches(4.45), Inches(5), Inches(0.5),
                today, font_size=12, font_color=RGBColor(120, 120, 120))
    add_footer_bar(slide)
    return slide

def make_toc_slide(prs, sections):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, COLOR_WHITE)
    add_top_bar(slide)
    add_textbox(slide, Inches(0.6), Inches(0.3), Inches(8), Inches(0.6),
                "汇报大纲", font_size=26, font_color=COLOR_IE_BLUE, bold=True)
    add_shape(slide, Inches(0.6), Inches(0.85), Inches(1.5), Inches(0.04),
              fill_color=COLOR_IE_ORANGE)
    items_per_col = (len(sections) + 1) // 2
    for idx, sec in enumerate(sections):
        col = idx // items_per_col
        row = idx % items_per_col
        x = Inches(0.8) + col * Inches(6.0)
        y = Inches(1.3) + row * Inches(0.48)
        c = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y + Inches(0.05),
                                   Inches(0.32), Inches(0.32))
        c.fill.solid(); c.fill.fore_color.rgb = COLOR_IE_BLUE
        c.line.fill.background()
        tf = c.text_frame; p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = str(idx + 1)
        r.font.size = Pt(10); r.font.color.rgb = COLOR_WHITE; r.font.bold = True
        add_textbox(slide, x + Inches(0.4), y, Inches(5), Inches(0.35),
                    sec, font_size=13, font_color=COLOR_DARK_GRAY)
    add_footer_bar(slide)
    return slide

def make_divider(prs, section_title, section_num=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, COLOR_IE_BLUE)
    add_shape(slide, Inches(0), Inches(3.0), SLIDE_W, Inches(0.06),
              fill_color=COLOR_IE_ORANGE)
    if section_num:
        add_textbox(slide, Inches(0.8), Inches(1.8), Inches(3), Inches(0.5),
                    f"Part {section_num}", font_size=16,
                    font_color=RGBColor(180, 210, 240))
    add_textbox(slide, Inches(0.8), Inches(3.3), Inches(11), Inches(1.2),
                section_title, font_size=30, font_color=COLOR_WHITE, bold=True)
    return slide

def make_content_slide(prs, title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, COLOR_WHITE)
    add_top_bar(slide)
    add_shape(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.75),
              fill_color=COLOR_BLUE_LIGHT_BG)
    add_shape(slide, Inches(0), Inches(0.75), SLIDE_W, Inches(0.04),
              fill_color=COLOR_IE_BLUE)
    add_textbox(slide, Inches(0.6), Inches(0.08), Inches(12), Inches(0.65),
                title, font_size=22, font_color=COLOR_IE_BLUE, bold=True)
    add_footer_bar(slide)
    return slide

def make_text_slide(prs, title, bullets, alert_text=None):
    """纯文字列表页 (支持小标题)"""
    slide = make_content_slide(prs, title)
    txBox = add_textbox(slide, Inches(0.8), Inches(1.05), Inches(11.5), Inches(5.5), "")
    tf = txBox.text_frame; tf.word_wrap = True
    for b in bullets:
        if isinstance(b, tuple):
            text, level, bold = b
            add_bullet(tf, text, level=level, bold=bold)
        elif isinstance(b, str) and b.startswith("##"):
            add_paragraph(tf, b[2:].strip(), font_size=14, font_color=COLOR_IE_BLUE, bold=True, space_before=Pt(10))
        else:
            add_bullet(tf, b)
    if alert_text:
        box = add_shape(slide, Inches(0.8), Inches(5.85), Inches(11.5), Inches(0.85),
                        fill_color=COLOR_RED_LIGHT_BG)
        box.line.color.rgb = COLOR_IE_RED; box.line.width = Pt(1)
        add_textbox(slide, Inches(1.0), Inches(5.95), Inches(11), Inches(0.65),
                    f"⚠ {alert_text}", font_size=12, font_color=COLOR_IE_RED)
    return slide

def make_two_col(prs, title, left_title, left_items, right_title, right_items,
                 left_alert=False, right_alert=False):
    slide = make_content_slide(prs, title)
    col_w = Inches(5.4); lx = Inches(0.8); rx = Inches(6.8)

    bg_l = COLOR_RED_LIGHT_BG if left_alert else COLOR_BLUE_LIGHT_BG
    fc_l = COLOR_IE_RED if left_alert else COLOR_IE_BLUE
    add_shape(slide, lx, Inches(1.1), col_w, Inches(0.45), fill_color=bg_l)
    add_textbox(slide, lx + Inches(0.15), Inches(1.12), col_w, Inches(0.4),
                left_title, font_size=13, font_color=fc_l, bold=True)
    txL = add_textbox(slide, lx + Inches(0.15), Inches(1.65), col_w - Inches(0.3), Inches(4.8), "")
    for item in left_items: add_bullet(txL.text_frame, item, font_size=12)

    bg_r = COLOR_GREEN_LIGHT_BG if right_alert else COLOR_BLUE_LIGHT_BG
    fc_r = COLOR_IE_GREEN if right_alert else COLOR_IE_BLUE
    add_shape(slide, rx, Inches(1.1), col_w, Inches(0.45), fill_color=bg_r)
    add_textbox(slide, rx + Inches(0.15), Inches(1.12), col_w, Inches(0.4),
                right_title, font_size=13, font_color=fc_r, bold=True)
    txR = add_textbox(slide, rx + Inches(0.15), Inches(1.65), col_w - Inches(0.3), Inches(4.8), "")
    for item in right_items: add_bullet(txR.text_frame, item, font_size=12)
    return slide

def make_product_slide(prs, title, specs, highlights=None, note=None):
    """产品规格页：左规格 + 右亮点"""
    slide = make_content_slide(prs, title)
    txBox = add_textbox(slide, Inches(0.8), Inches(1.1), Inches(7.5), Inches(5.5), "")
    tf = txBox.text_frame; tf.word_wrap = True
    for key, val in specs:
        p = tf.add_paragraph(); p.space_before = Pt(4); p.space_after = Pt(2)
        rk = p.add_run(); rk.text = f"▸ {key}："; rk.font.size = Pt(12)
        rk.font.color.rgb = COLOR_IE_BLUE; rk.font.bold = True; rk.font.name = "Microsoft YaHei"
        rv = p.add_run(); rv.text = val; rv.font.size = Pt(12)
        rv.font.color.rgb = COLOR_DARK_GRAY; rv.font.name = "Microsoft YaHei"
    if highlights:
        box = add_shape(slide, Inches(8.8), Inches(1.3), Inches(4.0), Inches(3.5),
                        fill_color=COLOR_BLUE_LIGHT_BG)
        box.line.color.rgb = COLOR_IE_BLUE; box.line.width = Pt(1.5)
        add_textbox(slide, Inches(9.0), Inches(1.4), Inches(3.6), Inches(0.4),
                    "💡 核心亮点", font_size=13, font_color=COLOR_IE_BLUE, bold=True)
        txH = add_textbox(slide, Inches(9.0), Inches(1.9), Inches(3.6), Inches(2.8), "")
        for h in highlights: add_bullet(txH.text_frame, h, bullet_char="✦", font_size=11)
    if note:
        nb = add_shape(slide, Inches(8.8), Inches(5.1), Inches(4.0), Inches(1.4),
                       fill_color=COLOR_RED_LIGHT_BG)
        nb.line.color.rgb = COLOR_IE_RED; nb.line.width = Pt(1)
        add_textbox(slide, Inches(9.0), Inches(5.2), Inches(3.6), Inches(1.2),
                    note, font_size=10, font_color=COLOR_IE_RED)
    return slide

def make_table_slide(prs, title, headers, rows, caption=None):
    slide = make_content_slide(prs, title)
    n_rows = len(rows) + 1; n_cols = len(headers)
    col_w = Inches(11.5 / n_cols); row_h = Inches(0.38)
    tbl_top = Inches(1.25); tbl_left = Inches(0.8)
    ts = slide.shapes.add_table(n_rows, n_cols, tbl_left, tbl_top,
                                col_w * n_cols, row_h * n_rows)
    tbl = ts.table
    for i in range(n_cols): tbl.columns[i].width = col_w
    for i, h in enumerate(headers):
        c = tbl.cell(0, i); c.text = h
        c.fill.solid(); c.fill.fore_color.rgb = COLOR_IE_BLUE
        for pg in c.text_frame.paragraphs:
            pg.alignment = PP_ALIGN.CENTER
            for rn in pg.runs: rn.font.size = Pt(10); rn.font.color.rgb = COLOR_WHITE; rn.font.bold = True
    for r, row in enumerate(rows):
        for c_i, txt in enumerate(row):
            c = tbl.cell(r + 1, c_i); c.text = str(txt)
            c.fill.solid()
            c.fill.fore_color.rgb = RGBColor(245, 248, 252) if r % 2 == 0 else COLOR_WHITE
            for pg in c.text_frame.paragraphs:
                pg.alignment = PP_ALIGN.CENTER
                for rn in pg.runs: rn.font.size = Pt(9); rn.font.color.rgb = COLOR_DARK_GRAY
    if caption:
        add_textbox(slide, Inches(0.8), tbl_top + row_h * n_rows + Inches(0.15),
                    Inches(11.5), Inches(0.4), caption, font_size=9,
                    font_color=RGBColor(120, 120, 120))
    return slide

def make_summary_slide(prs, title, summary_items, conclusion_text=None):
    slide = make_content_slide(prs, title)
    card_top = Inches(1.2)
    for i, item in enumerate(summary_items):
        x = Inches(0.6) + i * Inches(4.1)
        card = add_shape(slide, x, card_top, Inches(3.8), Inches(3.3),
                         fill_color=COLOR_BLUE_LIGHT_BG)
        card.line.color.rgb = COLOR_IE_BLUE; card.line.width = Pt(1)
        ns = add_shape(slide, x + Inches(0.15), card_top + Inches(0.15),
                       Inches(0.45), Inches(0.45), fill_color=COLOR_IE_BLUE)
        tf = ns.text_frame; p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = str(i + 1); r.font.size = Pt(14)
        r.font.color.rgb = COLOR_WHITE; r.font.bold = True
        add_textbox(slide, x + Inches(0.15), card_top + Inches(0.75),
                    Inches(3.5), Inches(2.4), item, font_size=10, font_color=COLOR_DARK_GRAY)
    if conclusion_text:
        cb = add_shape(slide, Inches(0.6), Inches(5.1), Inches(12.0), Inches(1.2),
                       fill_color=RGBColor(240, 245, 250))
        cb.line.color.rgb = COLOR_IE_BLUE; cb.line.width = Pt(1.5)
        add_textbox(slide, Inches(0.9), Inches(5.25), Inches(11.4), Inches(0.9),
                    f"📌 {conclusion_text}", font_size=13, font_color=COLOR_IE_BLUE, bold=True)
    return slide

def make_end_slide(prs, author="葛良晨"):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, COLOR_IE_BLUE)
    add_shape(slide, Inches(0), Inches(3.0), SLIDE_W, Inches(0.06), fill_color=COLOR_IE_ORANGE)
    add_textbox(slide, Inches(0), Inches(2.2), SLIDE_W, Inches(1.5),
                "谢谢！欢迎提问", font_size=36, font_color=COLOR_WHITE, bold=True,
                alignment=PP_ALIGN.CENTER)
    add_textbox(slide, Inches(0), Inches(3.8), SLIDE_W, Inches(0.8),
                author, font_size=20, font_color=RGBColor(180, 210, 240),
                alignment=PP_ALIGN.CENTER)
    add_textbox(slide, Inches(0), Inches(4.4), SLIDE_W, Inches(0.6),
                datetime.date.today().strftime("%Y年%m月%d日"), font_size=14,
                font_color=RGBColor(150, 190, 220), alignment=PP_ALIGN.CENTER)
    return slide


# ═══════════════════════ 主构建函数 ═══════════════════════

def build_merged_pptx(output_path: str):
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # ━━━━━━━ 标题页 ━━━━━━━
    make_title_slide(prs,
        title="AI 互联接口综合研究报告",
        subtitle="竞品分析 · 技术路线 · CXL 内存解耦 · 交换芯片 · 市场定位",
        author="葛良晨")

    # ━━━━━━━ 目录 ━━━━━━━
    sections = [
        "调研背景与动机",
        "结构性瓶颈与 Fabless 切入方向",
        "CXL 协议与内存解耦",
        "CXL 产品与生态 (Rambus / 澜起 / Marvell / XConn)",
        "物理层：Retimer 与 Smart Cable",
        "交换芯片：市场格局与国产突破",
        "竞争格局与全产品线评估",
        "供应链、生态与总结",
    ]
    make_toc_slide(prs, sections)

    # ====== Part 1: 调研背景与动机 ======
    make_divider(prs, "调研背景与动机", 1)

    make_text_slide(prs, "研究背景",
        bullets=[
            "AI 大模型训练/推理对互联带宽的需求呈指数级增长",
            "PCIe 6.0 (64 GT/s) 与 CXL 3.x 成为下一代互连标准",
            "GPU 峰值算力增速远超内存带宽与互联带宽增速",
            "国产替代浪潮下，国内厂商加速布局 Retimer / Switch / CXL 芯片",
            "##核心问题",
            "  • 数据传输能力不足导致计算单元有效利用率下降",
            "  • 系统瓶颈已从「计算受限」转向 Memory Wall / IO Wall / Power Wall",
            "  • 各细分赛道的技术壁垒与市场格局如何？",
            "  • 国产厂商的竞争定位与差异化优势在哪？",
        ],
        alert_text="AI 集群规模从千卡向万卡演进，互联瓶颈成为系统性能的关键制约因素")

    # ====== Part 2: 结构性瓶颈与 Fabless 切入方向 ======
    make_divider(prs, "结构性瓶颈与 Fabless 切入方向", 2)

    make_two_col(prs,
        title="结构性瓶颈与 Fabless 团队的进入空间",
        left_title="🔧 基本判断：数据移动开销已超过计算开销",
        left_items=[
            "GPU 峰值算力增速远超内存带宽与互联带宽增速",
            "系统瓶颈已从「计算受限」转向多维瓶颈",
            "Memory Wall / IO Wall / Power Wall",
            "Thermal Wall / Network Congestion / Rack-Level Sync",
            "这些瓶颈领域高度碎片化，大型厂商缺乏优先投入动力",
            "GPU 厂商关注毛利率更高的计算芯片，不愿深入 PHY 层",
        ],
        right_title="🌐 约束条件：避免与 NVIDIA 软件生态直接竞争",
        right_items=[
            "不涉足 AI GPU / 大模型训练芯片 / 通用推理 ASIC",
            "NVIDIA 的竞争壁垒不仅是硬件，更是整个生态：",
            "  CUDA + NCCL + TensorRT + PyTorch 集成",
            "  云平台适配 + 供应链控制的整体体系",
            "Fabless 团队应聚焦 NVIDIA 生态覆盖不到的细分瓶颈",
            "技术方向选择应从 Hyperscaler 的核心瓶颈出发",
        ],
        left_alert=True)

    make_table_slide(prs,
        title="五个技术方向总览",
        headers=["方向", "技术逻辑", "大型厂商投入不足原因", "核心壁垒"],
        rows=[
            ["Retimer / Smart Cable", "PCIe 6/7 PAM4 信号衰减严重", "毛利率偏低、PHY 层定位", "SI Validation + 系统兼容性"],
            ["CXL Memory Controller", "推理场景显存容量不足", "CXL 软件生态尚未成熟", "软件栈 + NUMA 调度优化"],
            ["Optical 控制 ASIC", "铜互联接近物理极限", "难点在 Mixed Signal + DSP", "PHY calibration + 研发周期长"],
            ["Telemetry ASIC", "集群可观测性不足", "硬件-系统软件交叉领域", "需与云厂商深度联调"],
            ["Power / VRM 控制", "单机架功耗 200-300kW", "传统 PMIC 未针对 AI 优化", "rack-level 系统理解"],
        ],
        caption="方向优先级: ① AI Retimer / Smart Cable > ② AI Infra Telemetry > ③ 专用 CXL Controller > ④ Optical ASIC > ⑤ Power Delivery")

    # ====== Part 3: CXL 协议与内存解耦 ======
    make_divider(prs, "CXL 协议与内存解耦", 3)

    make_two_col(prs,
        title="CXL：让内存「飞出」主板的开放标准",
        left_title="📡 CXL 是什么",
        left_items=[
            "Intel 牵头，AMD/ARM/NVIDIA/Google/Microsoft 共同支持的开放式互联标准",
            "底层构建在 PCIe 5.0/6.0 物理通道之上",
            "支持 CXL 的设备直接插在标准 PCIe 插槽",
            "核心创新：引入内存语义和缓存一致性协议",
            "CPU/GPU/加速器像访问本地内存一样直接访问彼此内存",
            "告别传统 PCIe 设备的多次数据拷贝 → 延迟大幅降低",
        ],
        right_title="⚡ 解决的终极问题",
        right_items=[
            "内存解耦：内存不再被「锁死」在主板 DIMM 插槽上",
            "可以「飞」到机架任何位置",
            "让多台服务器共享同一片物理内存",
            "传统架构：CPU → DDR 本地总线，插满即止",
            "CXL 架构：CPU → PCIe/CXL → CXL 内存设备 → DDR × N",
        ],
        right_alert=True)

    make_two_col(prs,
        title="CXL.mem vs RDMA：数据通路的本质差异",
        left_title="📦 RDMA：消息传递语义 (Message-based)",
        left_items=[
            "CPU → 内存 → 组装网络包 → NIC 网卡处理 QP & Doorbell",
            "→ 交换机 → NIC → 拆包 → 内存 → CPU",
            "延迟：1~5 μs",
            "多级拷贝 + 协议栈开销",
        ],
        right_title="⚡ CXL.mem：Load/Store 内存语义",
        right_items=[
            "CPU 执行 Load/Store 指令 → Cache Miss",
            "Home Agent → CXL Root Complex 路由",
            "→ CXL.mem Flit 硬件封装 (64B 缓存行)",
            "→ CXL Switch 硬件路由转发 → Type-3 设备执行",
            "→ 数据原路返回 CPU Cache",
            "延迟：100~300 ns — 零网卡、零拆包、零协议栈",
        ],
        left_alert=True, right_alert=True)

    make_text_slide(prs, "Leo：让 CXL 落地的智能内存控制器",
        bullets=[
            "##Leo 是什么",
            "Astera Labs 专为 CXL 标准研发的智能内存控制器芯片",
            "一端通过 CXL 协议接 PCIe 插槽，另一端接成排普通 DDR 内存条",
            "CPU「误以为」远端内存就是原生本地内存 → 超级转接头",
            "##痛点一：内存容量瓶颈 → E-Series 内存扩展",
            "LLM 推理/训练不缺算力，缺内存；模型参数动辄几百 GB",
            "Leo 让服务器通过 PCIe 插槽继续加内存，单台容量翻倍",
            "##痛点二：内存闲置孤岛 → P-Series 内存池化",
            "Server A (AI任务内存爆满) + Server B (70%闲置) → CXL 共享内存池",
            "按需划出，用完收回 → 消除 Stranded Memory，大幅降低 TCO",
        ])

    make_table_slide(prs,
        title="CXL 内存池化：业界代表案例",
        headers=["厂商", "产品/方案", "物理实现", "数据通路关键点"],
        rows=[
            ["Astera Labs", "Leo CXL 内存控制器 + Aries Retimer", "AEC 线缆 + E3.S 内存扩展模块", "Retimer CDR 恢复 PAM4 信号"],
            ["Samsung/SK Hynix", "CMM-D / CXL 内存模块", "E3.S 热插拔 + 自研 CXL 控制器 + DDR5", "CXL Flit 直接译为 DDR5 RAS/CAS"],
            ["UnifabriX", "CXL 3.0 Smart Memory Node", "2U 独立硬件 + CXL Switch + 大容量内存", "Switch 硬件路由 + Fabric Manager 动态重分配"],
            ["Ayar Labs", "TeraPHY 光学 I/O 芯粒 (CPO)", "硅光 Chiplet 与 CPU/Switch 同基板封装", "电信号 → 毫米内转光 → 单模光纤跨 50m"],
            ["阿里云", "磐久 CXL + LPO 直驱光模块", "LPO 去 DSP 光模块 + 自研 CXL 控制器", "CXL 时钟直接驱动光器件，省 DSP 延迟"],
        ],
        caption="层级划分: 机架内(铜): Astera/Samsung/SK Hynix 已量产 | 系统级池化: UnifabriX | 跨机柜(光): Ayar/阿里云指明终局")

    # ====== Part 4: CXL 产品与生态 ======
    make_divider(prs, "CXL 产品与生态 (Rambus / 澜起 / Marvell / XConn)", 4)

    make_product_slide(prs, "Rambus CXL 3.1 Controller IP",
        specs=[
            ("PHY", "PCIe 6.1 64 GT/s，Rambus 自研或用户自定义 PIPE 兼容 SerDes"),
            ("协议", "CXL 3.1 完整支持 (CXL.io / CXL.mem / CXL.cache)"),
            ("特点", "IP 核支持参数化配置，灵活适配不同应用场景"),
            ("验证", "提供基于兼容 FPGA SerDes 的参考设计，便于快速 FPGA 原型验证"),
        ],
        highlights=["完整 CXL 3.1 IP 解决方案", "参数化配置 + FPGA 快速原型", "降低芯片设计门槛"])

    make_product_slide(prs, "澜起 M88MX6852 — CXL 3.2 Type 3 内存扩展控制器 (MXC)",
        specs=[
            ("协议", "CXL 1.1 / 2.0 / 3.2，支持 CXL.mem 和 CXL.io"),
            ("速率", "PCIe 6.2 物理层，最高 64 GT/s (x8)"),
            ("内存", "DDR5 最高 DDR5-8000，支持 UDIMM / RDIMM / On-board DRAM"),
            ("集成", "内置 CXL 控制器 + DDR 控制器 + RISC-V 微处理器"),
            ("RAS", "丰富的 RAS 功能"),
            ("封装", "1211-ball FCBGA"),
        ],
        highlights=[
            "全球首款 MXC 芯片，CXL 3.1 送样",
            "部分节点全球领跑",
            "内存池化 + KV Cache pooling",
            "RISC-V 内置处理器灵活管理",
        ],
        note="应用场景: 内存 AIC 扩展卡、EDSFF 内存模组、内存数据库、云基础设施、AI/ML 负载")

    make_two_col(prs,
        title="Marvell Structera CXL 产品线",
        left_title="🔧 Structera A — 近存加速器",
        left_items=[
            "目标: 高带宽应用 (DLRM / ML / AI)",
            "PCIe 5.0 / CXL 2.0 x16",
            "4 通道 DDR5 6400 MT/s",
            "16× Arm Neoverse V2 @ 3.2 GHz",
            "带宽 200 GB/s，容量至 4 TB",
            "5nm 制程",
        ],
        right_title="📦 Structera X — 内存扩展控制器",
        right_items=[
            "目标: 高容量应用 (内存数据库等)",
            "PCIe 5.0 / CXL 2.0 1x16 或 2x8",
            "X2404: 4 通道 DDR4 3200, 12 条 DIMM",
            "X2504: 4 通道 DDR5 6400 MT/s",
            "支持双 CPU 同时访问",
            "内联 LZ4 压缩 + AES-XTS 256 加密",
        ])

    make_two_col(prs,
        title="XConn Technologies — CXL / PCIe 交换机芯片",
        left_title="XC50256 CXL 2.0 Switch",
        left_items=[
            "世界首款 CXL 2.0 / PCIe Gen 5.0 交换机芯片",
            "完全兼容 CXL 2.0/1.1 与 PCIe Gen 5.0",
            "支持 CXL.io / CXL.mem / CXL.cache",
            "CXL Fabric Manager + MLD + 多虚拟交换机 (VCS)",
            "总计 32 端口 (含 Bifurcation)，256 Lane",
            "交换容量 2,048 GB/s",
            "业内最低端口到端口延迟",
            "完整 RAS: ECC/Parity, DPC, Hot-Plug, Data Poisoning",
        ],
        right_title="XC51256 PCIe Gen 5.0 Switch",
        right_items=[
            "纯 PCIe 交换机芯片",
            "业内最低延迟 + 业界最大 Lane 数: 256 Lane",
            "x8/x16 端口 Bifurcation，最多 32 端口",
            "虚拟交换机 + 非透明桥 (NTB) + 多主机 IO 共享",
            "Hot-Plug / Hot-Add / Hot-Remove / Surprise-Plug",
            "内嵌逻辑分析仪 (ELA) + Debug GUI 诊断",
            "接口: I2C, SPI, UART, GPIO, JTAG, SRIS/SRNS",
        ],
        right_alert=True)

    # ====== Part 5: 物理层 Retimer & Smart Cable ======
    make_divider(prs, "物理层：Retimer 与 Smart Cable", 5)

    make_product_slide(prs, "澜起科技 M88RT61632 — 核心规格",
        specs=[
            ("通道数", "16 通道 PCIe 6.x / CXL 3.x Retimer"),
            ("最大速率", "64 GT/s（自研 PAM4 SerDes）"),
            ("链路预算", "高达 43 dB"),
            ("封装", "354-ball FCCSP 业界主流封装"),
            ("分叉模式", "1×16, 2×8, 4×4, 最多 8×2"),
            ("兼容性", "向下兼容 64/32/16/8/5/2.5 GT/s"),
            ("低功耗", "L0P / L1PM 低功耗模式"),
            ("管理接口", "SMBus / I3C 管理 + SPI/EEPROM 固件"),
        ],
        highlights=[
            "全球第 2 家推出 PCIe 6.0 Retimer",
            "自研 64 GT/s PAM4 SerDes IP",
            "链路预算 43 dB 反超 Astera (~36 dB)",
            "2025 年营收 54.56 亿元 (+49.9% YoY)",
        ])

    make_table_slide(prs,
        title="澜起 Retimer 产品线全貌",
        headers=["产品代", "状态", "关键指标", "市场地位"],
        rows=[
            ["PCIe 4.0 Retimer", "已量产", "—", "进入较晚，份额较小"],
            ["PCIe 5.0 / CXL 2.0", "已量产", "32 GT/s，出货超百万颗(2024)", "全球第2，市占率~10.9%"],
            ["PCIe 6.x / CXL 3.x", "送样中", "M88RT61632, 64 GT/s, 43 dB", "全球第2家推出"],
            ["PCIe 7.0 Retimer", "研发中", "128 GT/s 目标", "与 Astera 同步研发"],
            ["PCIe Switch", "研发中", "—", "国内空白填补"],
            ["以太网 PHY Retimer", "研发中", "—", "新赛道布局"],
        ],
        caption="📌 核心壁垒: 自研 64 GT/s PAM4 SerDes IP — 国内极少数掌握该核心技术的厂商")

    make_table_slide(prs,
        title="代表产品对比：澜起 vs Astera Labs vs 谱瑞",
        headers=["指标", "澜起 M88RT61632", "Astera Aries (PCIe 6)", "谱瑞 PS8926"],
        rows=[
            ["速率", "64 GT/s PAM4", "64 GT/s PAM4", "32 GT/s (PCIe 5.0)"],
            ["通道数", "16", "16", "16"],
            ["链路预算", "高达 43 dB", "~36 dB", "~36 dB"],
            ["协议", "PCIe 6.x / CXL 3.x", "PCIe 6.x / CXL 3.x", "PCIe 5.0 / CXL 2.0"],
            ["SerDes", "自研 PAM4 IP", "自研 DSP", "自研"],
            ["送样时间", "2025.01", "2024 (业界首家)", "已量产"],
            ["软件生态", "GUI + SDK", "COSMOS 遥测套件", "基础工具"],
        ],
        caption="📌 关键差距：Astera 先发 ~10 个月，COSMOS 软件生态构成最大护城河；澜起在链路预算 (43 dB) 上反超。")

    make_text_slide(prs, "AEC 产业链格局与速率演进",
        bullets=[
            "##产业链双层结构",
            "互连与线缆方案大厂: Credo (AEC先驱，主导HiWire联盟)、Amphenol/Molex (全球巨头)、FS/Approved Networks",
            "核心芯片供应商: Marvell/Broadcom (网络芯片双雄)、Astera Labs/Point2 (专注高速互连信号调理)",
            "##AEC vs DAC vs AOC 核心区别",
            "DAC (无源): 极低成本，但 800G 时距离压缩至 2.5m 内，无法跨机柜",
            "AEC (有源): 两端内嵌 Retimer 芯片 (CDR + 均衡)，功耗 6~12W，距离延伸至 7m",
            "AOC (有源光缆): 距离最长 (>100m)，成本最高，功耗最大",
            "##速率演进: 400G (56G PAM4, 7m) → 800G (112G PAM4, 5~7m) → 1.6T (224G PAM4, 2.5~3m)",
        ],
        alert_text="AEC 填补了短距铜缆 (DAC) 与长距光纤 (AOC) 之间的关键空白 — 400G/800G 已大规模商用，1.6T 进入量产初期")

    make_two_col(prs,
        title="Astera Labs Aries Smart Cable Module",
        left_title="🔧 产品定义与定位",
        left_items=[
            "Aries Retimer IC + 外围组件集成于 paddle card",
            "嵌入 AEC 线缆两端",
            "支持直连电缆 (Straight) 和扇出电缆 (Breakout)",
            "PCIe 5.0: 最远 7 米；PCIe 6.x: 最远 6 米",
            "支持 64/32/8/5/2.5 GT/s，自动链路均衡",
            "灵活分叉：1×16, 2×8, 4×4, 8×2",
        ],
        right_title="⚡ 高级特性",
        right_items=[
            "自动方向检测：对称电缆设计，即插即用",
            "自动极性校正 + Lane Reversal 支持",
            "支持热插拔 (Hot Plug / Hot Un-plug)",
            "接收端 Lane Margining (时序 + 电压)",
            "协议层 Loopback + PRBS 生成/检查",
            "不停机固件升级 (Non-disruptive FW Update)",
            "COSMOS 遥测套件：实时监控眼图张开度",
        ],
        right_alert=True)

    make_text_slide(prs, "Astera Labs Taurus — 以太网 AEC",
        bullets=[
            "##产品定义",
            "面向以太网 Scale-Out 场景的有源铜缆模块 (ACC)",
            "集成于 QSFP-DD / QSFP-DD800 / OSFP 标准接口",
            "替代粗重无源 DAC 和高成本 AOC，填充 3m 内成本最优区间",
            "支持 200G / 400G / 800G 三代速率",
            "##物理性能",
            "线速率：100G PAM4 / 50G PAM4 / 25G NRZ；最远 3 米",
            "端到端延迟 < 100 ns；功耗比 AOC 低高达 50%",
            "##高级特性",
            "每通道独立控制；深度线缆诊断 (眼图 + PRBS + loopback)",
            "安全固件加载 + 不停机固件升级；符合 CMIS 通用管理接口规范",
        ])

    make_table_slide(prs,
        title="Credo Semiconductor：全栈互联方案",
        headers=["产品线", "代表产品", "核心指标", "状态"],
        rows=[
            ["Blue Heron Retimer", "224G 多协议 Retimer", "3nm, UALink/ESUN/Ethernet, 40+ dB", "2026.01 送样, CQ3 量产"],
            ["ZeroFlap AEC", "有源电缆 (100G~1.6T)", "市占率 ~88%, 7m, 功耗降 50%", "已量产, PCIe 6 AEC 送样"],
            ["Cardinal DSP", "1.6T 光学 DSP", "3nm, 224G PAM4, EML+硅光驱动", "已送样"],
            ["PCIe Gen6 Retimer", "PCIe/CXL Retimer", "sub 7ns, 11W, 40 dB", "设计赢单, FY27 贡献"],
            ["OmniConnect", "112G VSR SerDes / Gearbox", "超短距, 低功耗", "FY28 量产"],
            ["ZeroFlap Optics", "800G 光收发模块", "PILOT 链路预测, 1000×可靠性", "2026.03 推出"],
        ],
        caption="财务: FY2025 $4.37亿 → FY2026E ~$9.85亿 (+125%) | 5大支柱: AEC+Optical DSP+Retimer+SerDes IP+Gearbox | TAM > $100亿")

    make_table_slide(prs,
        title="Astera Labs 产品矩阵与国内竞争对标",
        headers=["Astera 产品", "核心痛点", "国内对标企业", "竞争状态"],
        rows=[
            ["Aries (Retimer)", "物理层高速信号衰减", "澜起科技、谱瑞 (Parade)", "国内头部已量产 PCIe 5.0，研发 6.0"],
            ["Leo (CXL 内存控制器)", "内存容量扩展与池化", "澜起科技", "全球领跑阵营，CXL 3.1 送样"],
            ["Scorpio (交换芯片)", "多 GPU/节点复杂路由", "芯动科技、澜起科技(研发中)", "120 通道 PCIe 5.0 交换芯片已落地"],
            ["Taurus (智能 AEC)", "机架内线缆降本降功耗", "兆龙互连、瑞可达、澜起科技", "400G/800G AEC 制造成熟"],
        ],
        caption="⚠ 客观差距: ① 软件生态护城河 (COSMOS) — 管理数万条链路 | ② 超大规模实战检验 — Astera 在 AWS/微软万卡集群中历经考验")

    # ====== Part 6: 交换芯片 ======
    make_divider(prs, "交换芯片：市场格局与国产突破", 6)

    make_two_col(prs,
        title="PCIe Switch：LLM 万卡集群的「超级立交桥」",
        left_title="📊 市场格局：从垄断到国产突围",
        left_items=[
            "国际巨头: Broadcom + Microchip 长期垄断中高端",
            "Broadcom 在 AI 服务器占据绝对主导",
            "Astera Labs Scorpio 系列全球领先",
            "国产破局: 2025~2026 成为落地元年",
            "  PCIe 5.0 世代密集突破流片 + 商用部署",
        ],
        right_title="🎯 三大核心应用场景",
        right_items=[
            "GPU 高速互联: CPU/GPU/NIC/NVMe 间动态调度",
            "  消除带宽瓶颈，支持 All-Reduce 参数同步",
            "资源池化 (CDI): 物理分区/虚拟化灵活分配资源",
            "对等通信 (NTB): GPU Direct，绕过 CPU 直连",
        ],
        right_alert=True)

    make_table_slide(prs,
        title="国内代表企业对比",
        headers=["公司", "核心产品", "进展 (截至 2026)", "生态与合作"],
        rows=[
            ["芯动科技", "GX9120: 120 通道 PCIe 5.0", "2026.02 首发, 115ns 延迟, 16卡全互联", "国产主流服务器适配导入"],
            ["数渡科技", "104 通道 PCIe 5.0 Switch", "2026 初批量部署，联通星罗京津冀平台", "深度绑定华为昇腾 (910B/950)"],
            ["井芯微电子", "JXW8848 全国产自主 Switch", "2025 下半年量产，全链路自主可控", "AI 智算中心 + 高安全等级设备"],
            ["澜起科技", "PCIe 5.0/6.0 Retimer", "Retimer 细分赛道全球领先", "配合 PCIe Switch 解决信号完整性"],
        ],
        caption="万通发展 2025 下半年斥巨资控股数渡科技 — 资本市场对国产高速互联芯片赛道高度认可")

    make_product_slide(prs, "芯动科技 GX9120 — 全球首款 120 通道 PCIe Gen5 Switch",
        specs=[
            ("通道/端口", "120 条 PCIe Gen5 通道，30 个可灵活配置端口"),
            ("速率", "单通道 32 GT/s，全交叉非阻塞交换架构"),
            ("延迟", "超低延迟 115 ns（纳秒级）"),
            ("SerDes", "多级 FFE + 自适应 CTLE + 14 级 DFE，36 dB"),
            ("GPU 互联", "支持 16 卡 GPU 高效率互联"),
            ("模式", "多主机 / 端到端 / 非透明传输 (NTP)"),
            ("安全", "安全引擎 + PUF 物理不可克隆密钥 + 软件防火墙"),
            ("温度", "宽温 -40~+105°C, 全链路冗余纠错"),
        ],
        highlights=[
            "全球首款 120 通道 PCIe Gen5 Switch",
            "填补国产空白",
            "对标 Broadcom，填补 144/104 市场空白",
            "全自研底层 IP (MAC + SerDes PHY)",
            "已前瞻布局 PCIe 6.0 / CXL 3.0 Switch",
        ])

    make_table_slide(prs,
        title="代表产品对比：芯动 vs Astera vs Broadcom",
        headers=["指标", "芯动 GX9120", "Astera Scorpio P", "Broadcom (对标)"],
        rows=[
            ["通道数", "120", "32~320", "104 / 144"],
            ["端口数", "30 可配置", "全矩阵配置", "—"],
            ["协议", "PCIe Gen5", "PCIe 5.0 / CXL", "PCIe Gen5"],
            ["延迟", "115 ns", "超低延迟", "—"],
            ["SerDes", "14 级 DFE, 36 dB", "自研 DSP", "自研"],
            ["安全", "PUF + 防火墙", "Secure Boot + RoT", "—"],
            ["特色", "NTP 热插拔, 16 GPU", "Hypercast, In-Network Compute", "—"],
        ],
        caption="战略意义: 芯动扎根底层 IP 近 20 年，GX9120 打破博通/微芯长期垄断，已前瞻布局 144 通道 PCIe 6.0 / CXL 3.0 Switch")

    # ====== Part 7: 竞争格局与全产品线评估 ======
    make_divider(prs, "竞争格局与全产品线评估", 7)

    make_two_col(prs,
        title="技术壁垒与市场格局",
        left_title="🔧 硬件壁垒",
        left_items=[
            "SerDes IP: 64 GT/s PAM4 高速串行接口设计，国内仅澜起等极少数掌握",
            "信号完整性验证: 43 dB 链路预算下 DSP 均衡/DFE/CDR 综合调优",
            "互操作性认证: 跨 CPU/GPU/SSD/NIC 多厂商兼容性测试矩阵",
            "PCI-SIG 合规: PCIe 6.0 认证流程复杂，周期长",
        ],
        right_title="🌐 生态壁垒",
        right_items=[
            "COSMOS 软件: Astera 的遥测/诊断/OTA 套件是真正护城河",
            "客户绑定: Astera 已深度集成 NVIDIA/AWS/Meta 系统",
            "先发时间差: Astera 领先 ~10 个月量产 PCIe 6.0",
            "固件迭代: 大规模部署中的 bug 修复和性能调优需时间积累",
        ],
        left_alert=False, right_alert=True)

    make_table_slide(prs,
        title="全产品线竞争力矩阵",
        headers=["产品线", "技术成熟度", "市场空间", "壁垒高度", "团队匹配", "时机窗口", "综合评级"],
        rows=[
            ["Retimer / Smart Cable", "★★★★", "★★★★★", "★★★★", "★★★★", "★★★★★", "A"],
            ["Smart Switch", "★★★", "★★★★", "★★★★★", "★★★", "★★★★", "B+"],
            ["CXL 内存控制器", "★★★", "★★★★", "★★★★", "★★★★", "★★★★", "A-"],
            ["SuperNIC", "★★★", "★★★★", "★★★★★", "★★★", "★★★", "B"],
            ["D2D PHY IP", "★★★", "★★★★", "★★★★", "★★", "★★★", "B"],
            ["硅光 / CPO", "★★", "★★★★★", "★★★★★", "★★", "★★", "B+"],
        ])

    make_text_slide(prs, "澜起科技：财务表现与战略定位 (2025 年报)",
        bullets=[
            "##2025 年核心财务数据",
            "总营收 54.56 亿元 (+49.9% YoY)",
            "归母净利润 22.36 亿元 (+58.4%)",
            "互连类芯片收入 51.39 亿元 (+53.4%)，占 94.2%",
            "互连芯片毛利率 65.6%",
            "2026E 券商预测营收 75.07 亿 (+37.6%)，2027E 95.62 亿",
            "##战略卡位",
            "国内唯一在 Retimer + CXL MXC + AEC 全面对标 Astera 的厂商",
            "内存接口芯片 (RCD/DB/MRCD/MDB/CKD) 全球市占率 ~45%",
            "「内存互连 + 高速运力」双轮驱动",
            "CXL 3.1 MXC 芯片已送样，部分节点全球领跑",
        ],
        alert_text="与 Astera 的差距: Astera 2025 收入 $8.525 亿 (+115%)，净利润 $2.19 亿，市值 ~$230 亿。澜起市值 ~2200亿 RMB，Retimer 市占率差距显著")

    # ====== Part 8: 供应链、生态与总结 ======
    make_divider(prs, "供应链、生态与总结", 8)

    make_two_col(prs,
        title="关键供应链环节",
        left_title="🔧 上游依赖",
        left_items=[
            "SerDes IP: 核心壁垒，澜起自研掌握",
            "晶圆代工: 成熟制程 (16/12/7nm)，不依赖最先进节点",
            "封装: FCCSP 业界主流，不依赖 CoWoS 先进封装",
            "测试设备: Keysight/Tektronix/Viavi",
            "  PCIe 6.0 验证复杂度指数上升",
        ],
        right_title="🌐 下游客户",
        right_items=[
            "云厂商: 国内头部 CSP (阿里/腾讯/字节)",
            "服务器 OEM: 浪潮/新华三/联想",
            "GPU/AI 芯片厂商: 国产 GPU 互联需求",
            "线缆/模块厂商: AEC 有源线缆集成",
        ],
    )

    make_table_slide(prs,
        title="主要并购与融资事件",
        headers=["事件", "金额", "时间", "战略意图"],
        rows=[
            ["Marvell 收购 XConn", "$5.4 亿", "2026", "混合交换芯片 (Apollo 2)，打造开放 Fabric"],
            ["Marvell 收购 Celestial AI", "$32.5 亿", "2026", "光子 Fabric 技术，挑战 NVLink 生态"],
            ["Ayar Labs E 轮", "$5 亿", "2026.03", "CPO 光引擎量产，估值 $37.5 亿"],
            ["Enfabrica C 轮", "$1.15 亿", "2026", "SuperNIC ACF-S 芯片，3.2 Tbps"],
            ["Eliyan 战略融资", "$5000 万", "2025", "标准封装 D2D，AMD/ARM/Samsung 参投"],
        ],
        caption="2026 年 AI 互联领域并购融资空前活跃，Marvell 两笔重大收购合计 ~$38 亿")

    # ━━━━━ 总结页 ━━━━━
    make_summary_slide(prs,
        title="总结与建议",
        summary_items=[
            "PCIe Retimer 市场呈现双寡头格局：Astera (86%) 与澜起 (10.9%) 合计占 96.9%，国产替代空间巨大",
            "CXL 内存池化与 AI 集群互联是未来 3-5 年最高确定性增长赛道，澜起在 MXC 领域全球领跑",
            "交换芯片领域国产破局加速：芯动 GX9120 (120通道) 打破 Broadcom/Microchip 垄断，资本市场高度认可",
        ],
        conclusion_text="建议重点关注：澜起科技 (Retimer + CXL MXC + AEC 全面布局)、芯动科技 (交换芯片破局者)、以及 Credo/Astera 等海外龙头的技术演进方向。国内短板在软件生态 (COSMOS) 和超大规模实战验证，需持续投入。")

    # ━━━━━ 致谢页 ━━━━━
    make_end_slide(prs, author="葛良晨")

    # ━━━━━ 保存 ━━━━━
    prs.save(output_path)
    size = Path(output_path).stat().st_size
    print(f"✅ 合并版 PPTX 已生成: {output_path} ({size / 1024:.0f} KB)")
    print(f"   共 {len(prs.slides)} 页幻灯片")

# ═══════════════════════ 主入口 ═══════════════════════

def main():
    parser = argparse.ArgumentParser(
        description="AI 互联接口综合研究报告 — 合并版 PPTX 生成脚本")
    parser.add_argument("-o", "--output", default="interconnect_merged.pptx",
                        help="输出的 .pptx 文件路径 (默认: interconnect_merged.pptx)")
    args = parser.parse_args()
    build_merged_pptx(args.output)

if __name__ == "__main__":
    main()
